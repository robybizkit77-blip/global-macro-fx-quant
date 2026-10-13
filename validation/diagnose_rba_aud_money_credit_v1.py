#!/usr/bin/env python3
import io
import json
import urllib.request
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

RUNTIME = Path('live_data/sections/NATIVE_LIQ_DATA.json')
D3_URL = 'https://www.rba.gov.au/statistics/tables/xls/d03hist.xlsx'
D2_URL = 'https://www.rba.gov.au/statistics/tables/xls/d02hist.xlsx'
SOURCES = {
    'AUD.money': {'url': D3_URL, 'codes': ['DMABMS', 'DMABMN']},
    'AUD.credit': {'url': D2_URL, 'codes': ['DLCACBS', 'DLCACBN']},
}


def norm_month(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m')
    if isinstance(v, str):
        s = v.strip()
        for fmt in ('%Y-%m-%d', '%Y-%m', '%d/%m/%Y', '%b-%Y', '%b %Y'):
            try:
                return datetime.strptime(s, fmt).strftime('%Y-%m')
            except ValueError:
                pass
    return None


def fetch(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'GMFQ-validation/1.0'})
    with urllib.request.urlopen(req, timeout=60) as r:
        body = r.read()
    if len(body) < 10000:
        raise SystemExit(f'RBA workbook unexpectedly small: {url} bytes={len(body)}')
    return body


def find_code(ws, code):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        for cell in row:
            if str(cell.value).strip() == code:
                return cell.row, cell.column
    return None


def extract_column(ws, code_row, col):
    out = {}
    for r in range(code_row + 1, ws.max_row + 1):
        val = ws.cell(r, col).value
        if not isinstance(val, (int, float)):
            continue
        month = None
        for dc in range(1, col):
            month = norm_month(ws.cell(r, dc).value)
            if month:
                break
        if month:
            out[month] = float(val)
    return out


def extract_series(wb, code):
    for ws in wb.worksheets:
        hit = find_code(ws, code)
        if not hit:
            continue
        code_row, col = hit
        out = extract_column(ws, code_row, col)
        if out:
            header = [str(ws.cell(r, col).value or '').strip() for r in range(1, code_row)]
            return {'sheet': ws.title, 'code_row': code_row, 'column': col, 'header': header, 'values': out}
    return None


def runtime_series(rt, key):
    ccy, name = key.split('.')
    row = rt[ccy][name]
    out = {}
    for d, v in zip(row['dates'], row['values']):
        month = norm_month(d)
        if not month:
            raise SystemExit(f'unparseable runtime date for {key}: {d!r}')
        out[month] = float(v)
    return out


def compare(runtime, official):
    common = sorted(set(runtime) & set(official))
    exact = 0
    rounded_1dp = 0
    mismatches = []
    rounded_mismatches = []
    for m in common:
        rv, ov = runtime[m], official[m]
        if abs(rv - ov) <= 1e-6:
            exact += 1
        else:
            mismatches.append({'month': m, 'runtime': rv, 'official': ov, 'diff': rv - ov})
        if abs(rv - round(ov, 1)) <= 1e-9:
            rounded_1dp += 1
        else:
            rounded_mismatches.append({'month': m, 'runtime': rv, 'official': ov, 'official_1dp': round(ov, 1)})
    return {
        'overlap': len(common),
        'exact': exact,
        'exact_ratio': (exact / len(common)) if common else 0,
        'rounded_1dp_exact': rounded_1dp,
        'rounded_1dp_ratio': (rounded_1dp / len(common)) if common else 0,
        'mismatches_tail': mismatches[-8:],
        'rounded_mismatches_tail': rounded_mismatches[-12:],
        'first_common': common[0] if common else None,
        'last_common': common[-1] if common else None,
    }


def discover_d2_business(wb, runtime):
    rows = []
    for ws in wb.worksheets:
        # The RBA identifier row is discoverable from known legacy codes.
        known = find_code(ws, 'DLCACBS')
        if not known:
            continue
        code_row = known[0]
        for col in range(2, ws.max_column + 1):
            code = str(ws.cell(code_row, col).value or '').strip()
            if not code:
                continue
            header_parts = [str(ws.cell(r, col).value or '').strip() for r in range(1, code_row)]
            header = ' | '.join(x for x in header_parts if x)
            if 'business' not in header.lower():
                continue
            values = extract_column(ws, code_row, col)
            if not values:
                continue
            cmp = compare(runtime, values)
            rows.append({
                'code': code,
                'header': header,
                'latest_date': max(values),
                'latest_value': values[max(values)],
                **cmp,
            })
    rows.sort(key=lambda x: (x['rounded_1dp_exact'], x['overlap'], x['latest_date']), reverse=True)
    return rows[:30]


def main():
    rt = json.loads(RUNTIME.read_text(encoding='utf-8'))
    result = {'schema': 'GMFQ_RBA_AUD_MONEY_CREDIT_DIAGNOSTIC_V1', 'status': 'DIAGNOSTIC_ONLY', 'results': {}}
    workbooks = {}
    for key, spec in SOURCES.items():
        if spec['url'] not in workbooks:
            workbooks[spec['url']] = load_workbook(io.BytesIO(fetch(spec['url'])), data_only=True, read_only=False)
        wb = workbooks[spec['url']]
        runtime = runtime_series(rt, key)
        candidates = {}
        for code in spec['codes']:
            ex = extract_series(wb, code)
            if not ex:
                candidates[code] = {'found': False}
                continue
            candidates[code] = {
                'found': True,
                'sheet': ex['sheet'],
                'code_row': ex['code_row'],
                'column': ex['column'],
                'header': ' | '.join(x for x in ex['header'] if x),
                **compare(runtime, ex['values']),
                'runtime_last_date': max(runtime),
                'runtime_last_value': runtime[max(runtime)],
                'official_on_runtime_last_date': ex['values'].get(max(runtime)),
                'official_latest_date': max(ex['values']),
                'official_latest_value': ex['values'][max(ex['values'])],
            }
        row = {'source_url': spec['url'], 'candidates': candidates}
        if key == 'AUD.credit':
            row['business_series_discovery'] = discover_d2_business(wb, runtime)
        result['results'][key] = row
    print(json.dumps(result, indent=2, ensure_ascii=False))

if __name__ == '__main__':
    main()
