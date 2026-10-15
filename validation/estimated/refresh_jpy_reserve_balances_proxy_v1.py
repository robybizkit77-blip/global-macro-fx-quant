#!/usr/bin/env python3
from __future__ import annotations

import io
import json
import math
import re
import urllib.request
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

OUT = Path("live_data/sections/ESTIMATED_RESIDUAL_GAPS.json")
WORKBOOK_URL = "https://www.boj.or.jp/en/statistics/boj/other/mb/mblong.xlsx"
KNOWN_AUG_2026 = 3_905_362.0
KNOWN_SEP_2026 = 3_849_859.0


def _fetch_workbook() -> bytes:
    req = urllib.request.Request(
        WORKBOOK_URL,
        headers={
            "User-Agent": "Mozilla/5.0 (GMFQ BoJ official reserve-balances proxy refresh)",
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
        },
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        if resp.status != 200:
            raise RuntimeError(f"BoJ workbook HTTP {resp.status}")
        data = resp.read()
    if len(data) < 10_000 or data[:2] != b"PK":
        raise RuntimeError("BoJ workbook response is not a valid XLSX payload")
    return data


def _number(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        x = float(value)
        return x if math.isfinite(x) else None
    if isinstance(value, str):
        s = value.strip().replace(",", "").replace("−", "-")
        if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", s):
            try:
                x = float(s)
                return x if math.isfinite(x) else None
            except ValueError:
                return None
    return None


def _period(value):
    if isinstance(value, datetime):
        return f"{value.year:04d}-{value.month:02d}"
    if isinstance(value, date):
        return f"{value.year:04d}-{value.month:02d}"
    if value is None:
        return None
    s = str(value).strip()
    m = re.search(r"(20\d{2})\s*[/\-.年]\s*(1[0-2]|0?[1-9])(?:\s*月)?", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    return None


def _row_period(row_values):
    for value in row_values[:12]:
        p = _period(value)
        if p:
            return p
    return None


def _extract_series(workbook_bytes: bytes):
    wb = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    candidates = []

    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        max_cols = max(len(r) for r in rows)

        for anchor_r, row in enumerate(rows):
            row_period = _row_period(row)
            if row_period != "2026-08":
                continue
            for anchor_c, cell in enumerate(row):
                num = _number(cell)
                if num is None or abs(num - KNOWN_AUG_2026) > 0.5:
                    continue

                series = {}
                for r in rows:
                    if anchor_c >= len(r):
                        continue
                    p = _row_period(r)
                    v = _number(r[anchor_c])
                    if p and v is not None and v > 0:
                        series[p] = v

                if abs(series.get("2026-08", -1) - KNOWN_AUG_2026) > 0.5:
                    continue
                if "2026-09" in series and abs(series["2026-09"] - KNOWN_SEP_2026) > 0.5:
                    raise RuntimeError(
                        f"BoJ September anchor mismatch on {ws.title}: {series['2026-09']}"
                    )

                header_text = []
                for rr in range(max(0, anchor_r - 12), min(len(rows), anchor_r + 1)):
                    if anchor_c < len(rows[rr]) and rows[rr][anchor_c] is not None:
                        header_text.append(str(rows[rr][anchor_c]))
                candidates.append({
                    "sheet": ws.title,
                    "column": anchor_c + 1,
                    "header_context": " | ".join(header_text[-8:]),
                    "series": series,
                })

    if not candidates:
        raise RuntimeError("Could not locate the certified Aug-2026 Reserve Balances anchor in BoJ workbook")

    candidates.sort(key=lambda c: max(c["series"]), reverse=True)
    chosen = candidates[0]
    if len(candidates) > 1:
        best_latest = max(chosen["series"])
        equally_current = [c for c in candidates if max(c["series"]) == best_latest]
        if len(equally_current) > 1:
            semantic = [c for c in equally_current if "reserve" in c["header_context"].lower()]
            if len(semantic) == 1:
                chosen = semantic[0]
            elif len(semantic) > 1:
                chosen = semantic[0]

    if "2026-09" not in chosen["series"]:
        raise RuntimeError("BoJ workbook has not exposed the already-published September 2026 reserve balance")
    if abs(chosen["series"]["2026-09"] - KNOWN_SEP_2026) > 0.5:
        raise RuntimeError("Selected BoJ series failed September 2026 continuity anchor")
    return chosen


def _pct(a, b):
    if b in (None, 0):
        return None
    return round((a / b - 1.0) * 100.0, 4)


def _previous_month(period: str, months: int) -> str:
    y, m = map(int, period.split("-"))
    idx = y * 12 + (m - 1) - months
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def main():
    raw = _fetch_workbook()
    chosen = _extract_series(raw)
    series = chosen["series"]
    latest_period = max(series)
    latest = series[latest_period]
    prev = series.get(_previous_month(latest_period, 1))
    prev3 = series.get(_previous_month(latest_period, 3))

    if not (1_000_000 <= latest <= 8_000_000):
        raise RuntimeError(f"JPY reserve balance outside plausibility band: {latest}")

    doc = json.loads(OUT.read_text(encoding="utf-8"))
    if doc.get("official_lanes_untouched") is not True:
        raise RuntimeError("Residual governance is not fail-closed")
    lane = doc["lanes"]["JPY.reserve_balances"]

    original_runtime = lane.get("official_runtime_reference", 3_905_822.0)
    lane.pop("official_runtime_reference", None)
    lane.pop("absolute_difference", None)
    lane.pop("difference_pct", None)
    lane.update({
        "status": "ESTIMATED_ACTIVE_AUTO_REFRESH_OFFICIAL_PROXY",
        "confidence": "A/B",
        "proxy_as_of": latest_period,
        "proxy_value": round(latest, 6),
        "change_1m_pct": _pct(latest, prev),
        "change_3m_pct": _pct(latest, prev3),
        "source": "Bank of Japan official Monetary Base long-term workbook · Reserve Balances",
        "source_endpoint": WORKBOOK_URL,
        "source_sheet": chosen["sheet"],
        "source_column": chosen["column"],
        "source_header_context": chosen["header_context"],
        "fetched_at_utc": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "official_canonical_untouched": True,
        "original_runtime_validation": {
            "period": "2026-08",
            "official_proxy_value": KNOWN_AUG_2026,
            "runtime_reference": original_runtime,
            "absolute_difference": round(abs(KNOWN_AUG_2026 - original_runtime), 6),
            "difference_pct": round(abs(KNOWN_AUG_2026 - original_runtime) / original_runtime * 100.0, 4),
            "original_runtime_vintage_equivalence_claimed": False,
        },
        "continuity_anchor": {
            "period": "2026-09",
            "expected_official_value": KNOWN_SEP_2026,
            "matched": True,
        },
        "interpretation": "Semantically exact BoJ Reserve Balances series, auto-refreshed from the official workbook. Suitable for live level/direction/momentum; original runtime vintage equivalence remains explicitly unclaimed."
    })

    OUT.write_text(json.dumps(doc, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": "PASS",
        "lane": "JPY.reserve_balances",
        "proxy_status": lane["status"],
        "proxy_as_of": latest_period,
        "proxy_value": lane["proxy_value"],
        "change_1m_pct": lane["change_1m_pct"],
        "change_3m_pct": lane["change_3m_pct"],
        "source_sheet": chosen["sheet"],
        "source_column": chosen["column"],
        "aug_2026_anchor": series["2026-08"],
        "sep_2026_anchor": series["2026-09"],
        "official_canonical_untouched": True,
    }))


if __name__ == "__main__":
    main()
