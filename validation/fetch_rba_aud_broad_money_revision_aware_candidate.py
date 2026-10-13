#!/usr/bin/env python3
from __future__ import annotations

import argparse
import io
import json
import urllib.request
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_RUNTIME = Path('live_data/sections/NATIVE_LIQ_DATA.json')
SOURCE_URL = 'https://www.rba.gov.au/statistics/tables/xls/d03hist.xlsx'
SERIES_KEY = 'DMABMS'
TOL = 1e-6


def norm_month(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m')
    if isinstance(v, str):
        s=v.strip()
        for fmt in ('%Y-%m-%d','%Y-%m','%d/%m/%Y','%b-%Y','%b %Y'):
            try:
                return datetime.strptime(s,fmt).strftime('%Y-%m')
            except ValueError:
                pass
    return None


def next_month(v):
    y,m=map(int,v.split('-'))
    return f'{y+1:04d}-01' if m==12 else f'{y:04d}-{m+1:02d}'


def fetch_official():
    req=urllib.request.Request(SOURCE_URL,headers={'User-Agent':'GMFQ-validation/1.0'})
    with urllib.request.urlopen(req,timeout=60) as r:
        body=r.read()
    if len(body)<10000:
        raise SystemExit(f'RBA workbook unexpectedly small: {len(body)}')
    wb=load_workbook(io.BytesIO(body),data_only=True,read_only=False)
    for ws in wb.worksheets:
        hit=None
        for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,40)):
            for cell in row:
                if str(cell.value).strip()==SERIES_KEY:
                    hit=(cell.row,cell.column)
                    break
            if hit: break
        if not hit: continue
        code_row,col=hit
        header=[str(ws.cell(r,col).value or '').strip() for r in range(1,code_row)]
        out={}
        for rr in range(code_row+1,ws.max_row+1):
            val=ws.cell(rr,col).value
            if not isinstance(val,(int,float)): continue
            month=None
            for dc in range(1,col):
                month=norm_month(ws.cell(rr,dc).value)
                if month: break
            if month: out[month]=float(val)
        if out:
            return {'sheet':ws.title,'header':' | '.join(x for x in header if x)},out
    raise SystemExit(f'RBA series {SERIES_KEY} not found')


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--runtime',default=str(DEFAULT_RUNTIME))
    ap.add_argument('--out')
    args=ap.parse_args()
    payload=json.loads(Path(args.runtime).read_text(encoding='utf-8'))
    lane=payload['AUD']['money']
    dates=[norm_month(x) for x in lane.get('dates') or []]
    vals=[float(x) for x in lane.get('values') or []]
    if not dates or None in dates or len(dates)!=len(vals): raise SystemExit('invalid AUD.money runtime')
    if dates[-1]!=norm_month(lane.get('last_date')) or abs(vals[-1]-float(lane.get('last_value')))>TOL:
        raise SystemExit('AUD.money runtime tail inconsistency')
    runtime=dict(zip(dates,vals))
    meta,official=fetch_official()
    common=sorted(set(runtime)&set(official))
    if len(common)<60: raise SystemExit(f'insufficient AUD.money overlap: {len(common)}')
    exact=[]; rounded=[]; mismatches=[]
    for dt in common:
        rv,ov=runtime[dt],official[dt]
        if abs(rv-ov)<=TOL: exact.append(dt)
        if abs(rv-round(ov,1))<=1e-9: rounded.append(dt)
        else: mismatches.append({'date':dt,'runtime':rv,'official':ov,'official_1dp':round(ov,1),'diff':rv-ov})
    last=dates[-1]
    if last not in official: raise SystemExit('RBA current vintage missing runtime last month')
    new_dates=sorted(dt for dt in official if dt>last)
    expected=next_month(last); new=[]
    for dt in new_dates:
        if dt!=expected: raise SystemExit(f'non-contiguous AUD.money month: got {dt}, expected {expected}')
        new.append({'date':dt,'value':round(official[dt],1)})
        expected=next_month(dt)
    latest=max(official)
    result={
      'schema':'GMFQ_LIQUIDITY_AUD_BROAD_MONEY_REVISION_AWARE_CANDIDATE_V1',
      'status':'GREEN_SOURCE_CANDIDATE_ONLY',
      'source':{'publisher':'Reserve Bank of Australia','dataset':'D3 Monetary Aggregates','series_key':SERIES_KEY,'header':meta['header'],'sheet':meta['sheet'],'frequency':'Monthly','url':SOURCE_URL},
      'runtime':{'lane':'AUD.money','count':len(dates),'last_date':last,'last_value':vals[-1]},
      'proof':{'revision_aware_append_only':True,'historical_runtime_values_frozen':True,'new_observation_transform':'ROUND_1DP','overlap_count':len(common),'exact_match_count':len(exact),'rounded_1dp_match_count':len(rounded),'rounded_1dp_match_ratio':len(rounded)/len(common),'current_vintage_mismatch_count_after_1dp':len(mismatches),'mismatches_tail':mismatches[-12:],'runtime_last_official_current_vintage':official[last],'runtime_last_official_1dp':round(official[last],1),'runtime_last_current_vintage_diff':vals[-1]-round(official[last],1)},
      'latest_date':latest,'latest_raw_value':official[latest],'latest_value_1dp':round(official[latest],1),'new_observations':new,'candidate_update_available':bool(new),'candidate_pipeline_ready':True,'runtime_mutated':False,'publication_performed':False}
    text=json.dumps(result,indent=2,ensure_ascii=False)
    if args.out: Path(args.out).write_text(text+'\n',encoding='utf-8')
    else: print(text)

if __name__=='__main__': main()
