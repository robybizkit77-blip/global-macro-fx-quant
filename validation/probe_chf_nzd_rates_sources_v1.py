#!/usr/bin/env python3
"""Non-mutating access/shape probe of official SNB and RBNZ government yield sources.

No secondary providers, no data publication and no changes to NATIVE_RATES_DATA.
If source access, tenor identity, or common observation date is unclear, WITHHELD.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import re
import ssl
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

SNB_LAST_UPDATE = "https://data.snb.ch/api/cube/rendeiduebd/lastUpdate"
SNB_CSV = "https://data.snb.ch/api/cube/rendeiduebd/data/csv/en"
RBNZ_B2 = "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily-close.xlsx"
RBNZ_B2_ALTERNATE = "https://rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily-close.xlsx"
RBNZ_B2_PUBLIC_PAGE = "https://www.rbnz.govt.nz/statistics/series/exchange-and-interest-rates/wholesale-interest-rates"
DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
TERMS = {"2J": "2Y", "10J": "10Y"}
UA = "GlobalMacroFXQuant/1.0 (+read-only official source verification)"


def fetch(url: str, *, accept: str, limit: int = 3_000_000) -> tuple[bytes, dict]:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": accept})
    with urllib.request.urlopen(req, timeout=22, context=ssl.create_default_context()) as response:
        raw = response.read(limit + 1)
        if len(raw) > limit:
            raise ValueError("response larger than probe limit")
        return raw, {
            "http": response.status,
            "content_type": response.headers.get("content-type", ""),
            "bytes": len(raw),
        }


def parse_snb_csv(payload: bytes) -> dict:
    """Only publish evidence for paired 2J/10J CHF rows with unambiguous dimension codes."""
    text = payload.decode("utf-8-sig")
    first = text[:3000]
    delimiter = ";" if first.count(";") >= first.count(",") else ","
    rows = csv.reader(io.StringIO(text), delimiter=delimiter)
    found: dict[str, dict[str, float]] = {}
    candidates = 0
    for row in rows:
        if len(row) < 4:
            continue
        cells = [x.strip().strip('"') for x in row]
        dates = [x for x in cells if DATE.fullmatch(x)]
        if len(dates) != 1 or cells.count("CHF") != 1:
            continue
        tenors = [v for v in TERMS if cells.count(v) == 1]
        if len(tenors) != 1:
            continue
        # Require the numeric value in the final column; do not infer it from other data.
        try:
            value = float(cells[-1].replace(",", "."))
            dt.date.fromisoformat(dates[0])
        except (ValueError, OverflowError):
            continue
        if not (-3.0 < value < 18.0):
            continue
        candidates += 1
        tenor = TERMS[tenors[0]]
        rowdate = dates[0]
        day = found.setdefault(rowdate, {})
        if tenor in day and day[tenor] != value:
            raise ValueError("conflicting duplicate SNB same-date observation")
        day[tenor] = value
    paired = sorted((k, v) for k, v in found.items() if set(v) == {"2Y", "10Y"})
    if not paired:
        raise ValueError(f"no unambiguous same-date CHF 2Y/10Y pair; candidate_rows={candidates}")
    date, values = paired[-1]
    return {"as_of": date, "2Y": values["2Y"], "10Y": values["10Y"],
            "paired_days": len(paired), "recognized_rows": candidates,
            "basis": "SNB rendeiduebd spot Swiss Confederation CHF 2J/10J"}


def freshness(iso: str, *, max_age: int = 5) -> str:
    delta = (dt.datetime.now(dt.timezone.utc).date() - dt.date.fromisoformat(iso)).days
    return "CURRENT" if 0 <= delta <= max_age else "WITHHELD_STALE_OR_FUTURE"


def error_summary(exc: Exception) -> dict:
    status = getattr(exc, "code", None)
    if isinstance(exc, urllib.error.HTTPError):
        return {"kind": "HTTP", "status": status, "reason": str(exc.reason)[:100]}
    return {"kind": type(exc).__name__, "reason": str(exc)[:180]}


def probe_snb() -> dict:
    report = {"authority": "Swiss National Bank", "source": SNB_CSV,
              "status": "WITHHELD", "writes_operational_rates": False}
    try:
        raw, http = fetch(SNB_LAST_UPDATE, accept="application/json,text/plain")
        report["last_update_endpoint"] = {"http": http["http"],
                                            "bytes": http["bytes"]}
    except Exception as exc:
        report["last_update_endpoint"] = {"error": error_summary(exc)}
    params = urllib.parse.urlencode({"dimSel": "D0(CHF),D1(2J,10J)",
                                    "fromDate": "2026-09-01"})
    try:
        raw, http = fetch(SNB_CSV + "?" + params, accept="text/csv,text/plain")
        report["access"] = http
        result = parse_snb_csv(raw)
        report["paired_observations"] = result
        report["status"] = "OFFICIAL_PAIRED_EVIDENCE_" + freshness(result["as_of"])
    except Exception as exc:
        report["error"] = error_summary(exc)
    return report


def probe_rbnz() -> dict:
    report = {"authority": "Reserve Bank of New Zealand", "source": RBNZ_B2,
              "status": "WITHHELD", "writes_operational_rates": False}
    try:
        raw, http = fetch(RBNZ_B2,
                          accept="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream",
                          limit=10_000_000)
        report["access"] = http
        if not raw.startswith(b"PK\x03\x04") or len(raw) < 25_000:
            raise ValueError("official B2 response not a credible XLSX workbook")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            report["workbook"] = {
                "sheets": workbook.sheetnames[:12],
                "size": len(raw),
            }
            # Workbook access alone is not enough to assert a 2Y/10Y matched close.
            report["status"] = "OFFICIAL_WORKBOOK_ACCESSIBLE_NOT_CERTIFIED"
        finally:
            workbook.close()
    except Exception as exc:
        report["error"] = error_summary(exc)
    # Verify the official publication's own second hostname and published
    # page, without evading robots, authentication or blocking controls.
    # These checks never populate operational yields.
    report["official_access_routes"] = {}
    for label, url, accept, limit in (
        ("canonical_xlsx", RBNZ_B2,
         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 10_000_000),
        ("official_non_www_xlsx", RBNZ_B2_ALTERNATE,
         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 10_000_000),
        ("official_publication_page", RBNZ_B2_PUBLIC_PAGE, "text/html", 2_000_000),
    ):
        if label=="canonical_xlsx":
            report["official_access_routes"][label] = (
                report.get("access") or {"error": report.get("error", "unavailable")})
            continue
        try:
            payload, meta = fetch(url, accept=accept, limit=limit)
            if label.endswith("_xlsx"):
                if not payload.startswith(b"PK\x03\x04"):
                    raise ValueError("not an XLSX signature")
                meta["xlsx_signature"] = "PASS"
                meta["sha256"] = __import__("hashlib").sha256(payload).hexdigest()
            else:
                meta["contains_B2"] = b"Wholesale interest rates" in payload or b"Wholesale Interest Rates" in payload
            report["official_access_routes"][label] = meta
            if label.endswith("_xlsx") and report["status"]=="WITHHELD":
                report["status"] = "ALTERNATE_OFFICIAL_XLSX_ACCESSIBLE_NOT_CERTIFIED"
        except Exception as exc:
            report["official_access_routes"][label] = {"error": error_summary(exc)}
    return report


def self_test() -> None:
    csv_text = ("Date;D0;D1;Value\n"
                "2026-10-15;CHF;2J;0.31\n"
                "2026-10-15;CHF;10J;0.58\n"
                "2026-10-16;CHF;2J;0.32\n"
                "2026-10-16;CHF;10J;0.62\n").encode()
    assert parse_snb_csv(csv_text)["as_of"] == "2026-10-16"
    assert parse_snb_csv(csv_text)["10Y"] == 0.62
    for bad in [b"Date;D0;D1;Value\n2026-10-16;CHF;2J;0.32\n",
                b"Date;D0;D1;Value\n2026-10-16;CHF;2J;0.32\n2026-10-15;CHF;10J;0.62\n",
                b"Date;D0;D1;Value\n2026-10-16;CHF;2J;0.32\n2026-10-16;CHF;2J;0.77\n2026-10-16;CHF;10J;0.62\n"]:
        try:
            parse_snb_csv(bad)
            raise AssertionError("bad SNB official pair accepted")
        except ValueError:
            pass
    print("CHF_NZD_OFFICIAL_ACCESS_PROBE_SELFTEST_GREEN")


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--self-test", action="store_true")
    p.add_argument("--output", default="")
    args = p.parse_args()
    if args.self_test:
        self_test()
        return
    report = {
        "schema": "GMFQ_CHF_NZD_OFFICIAL_ACCESS_PROBE_V1",
        "generated_at_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "mode": "READ_ONLY_DIAGNOSTIC",
        "policy": "Never promote from access-only; require a separately verified paired-date candidate.",
        "sources": {"CHF": probe_snb(), "NZD": probe_rbnz()},
    }
    encoded = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        target = Path(args.output)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(encoded, encoding="utf-8")
    print(encoded)


if __name__ == "__main__":
    main()
