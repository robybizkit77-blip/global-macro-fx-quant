#!/usr/bin/env python3
from __future__ import annotations
import io,json,re,urllib.parse,urllib.request
from html.parser import HTMLParser
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
EVID=ROOT/'validation/NZD_RBNZ_2Y_PROBE_V1_2026-10-07.json'
PAGE='https://www.rbnz.govt.nz/statistics/series/exchange-and-interest-rates/wholesale-interest-rates'
UA={'User-Agent':'Mozilla/5.0 (compatible; global-macro-fx-quant/1.0)'}

class L(HTMLParser):
    def __init__(self): super().__init__(); self.links=[]; self._href=None; self._txt=[]
    def handle_starttag(self,tag,attrs):
        if tag=='a': self._href=dict(attrs).get('href'); self._txt=[]
    def handle_data(self,d):
        if self._href is not None:self._txt.append(d)
    def handle_endtag(self,tag):
        if tag=='a' and self._href is not None:
            self.links.append((self._href,' '.join(self._txt).strip())); self._href=None; self._txt=[]

def get(url):
    req=urllib.request.Request(url,headers=UA)
    with urllib.request.urlopen(req,timeout=60) as r:return r.read()

def main():
    raw=get(PAGE).decode('utf-8',errors='replace'); p=L();p.feed(raw)
    candidates=[]
    for href,txt in p.links:
        s=re.sub(r'\s+',' ',txt).lower()
        if 'wholesale interest rates' in s and 'daily close' in s and '2018-current' in s and ('xlsx' in s or href.lower().endswith('.xlsx')):
            candidates.append(urllib.parse.urljoin(PAGE,href))
    if len(set(candidates))!=1:raise RuntimeError(f'expected one RBNZ B2 daily-close 2018-current link; got {candidates}')
    url=candidates[0]; data=get(url)
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    matches=[]
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,30),values_only=True):
            vals=[str(x).strip() if x is not None else '' for x in row]
            for i,v in enumerate(vals):
                if re.search(r'2\s*year',v,re.I) and ('government' in ' '.join(vals).lower() or 'secondary market' in ' '.join(vals).lower()):
                    matches.append({'sheet':ws.title,'column_index':i+1,'header':v,'row':vals})
    payload={'schema':'GMFQ_NZD_RBNZ_2Y_PROBE_V1','status':'PASS_PROBE_NOT_CERTIFIED' if matches else 'FAIL','created_at':'2026-10-07','source':'Reserve Bank of New Zealand Wholesale interest rates (B2)','page_url':PAGE,'xlsx_url':url,'workbook_sheets':wb.sheetnames,'header_matches':matches[:20],'note':'Probe only. Certification requires exact column identity, daily date parsing, anchors and same-basis coverage gate.','changes_engine_rules':False,'changes_live_data':False,'changes_oos_baseline':False}
    EVID.write_text(json.dumps(payload,indent=2)+'\n',encoding='utf-8');print(json.dumps(payload,indent=2))
    if not matches:raise SystemExit(1)
if __name__=='__main__':main()
