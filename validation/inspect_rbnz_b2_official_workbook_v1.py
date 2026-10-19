#!/usr/bin/env python3
"""Read-only evidence inventory of an official RBNZ B2 XLSX supplied by an authorized user.

Never guess the 2Y/10Y mapping or write any operational data.
"""
from __future__ import annotations
import argparse, hashlib, io, json, re
from pathlib import Path
from openpyxl import load_workbook

OFFICIAL_NAME="hb2-daily-close.xlsx"

def inspect(path: Path) -> dict:
    raw=path.read_bytes()
    if path.name.lower()!=OFFICIAL_NAME:
        raise ValueError("Official filename must be hb2-daily-close.xlsx")
    if len(raw)<20000 or not raw.startswith(b"PK\x03\x04"):
        raise ValueError("Invalid official XLSX byte signature or implausible size")
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    try:
        evidence={"schema":"GMFQ_RBNZ_B2_MAPPING_DISCOVERY_V1",
          "status":"REVIEW_REQUIRED_NOT_CERTIFIED",
          "official_original_filename":OFFICIAL_NAME,
          "source_url":"https://rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily-close.xlsx",
          "sha256":hashlib.sha256(raw).hexdigest(),
          "size":len(raw),"worksheets":[],"writes_live_data":False}
        for sheet in wb:
            headings=[]
            for idx,row in enumerate(sheet.iter_rows(min_row=1,max_row=min(35,sheet.max_row),max_col=min(90,sheet.max_column),values_only=True),start=1):
                hits=[]
                for col,val in enumerate(row,start=1):
                    if val is None:continue
                    txt=re.sub(r"\s+"," ",str(val).strip())
                    if not txt:continue
                    if any(k in txt.lower() for k in ("date","government bond","closing yield","2 year","10 year","bond yields")):
                        hits.append({"col":col,"text":txt[:140]})
                if hits:headings.append({"row":idx,"cells":hits})
            evidence["worksheets"].append({"sheet":sheet.title,"rows":sheet.max_row,"columns":sheet.max_column,"headings":headings})
        return evidence
    finally:
        wb.close()

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("workbook",type=Path)
    parser.add_argument("--output",type=Path)
    args=parser.parse_args()
    result=inspect(args.workbook)
    rendered=json.dumps(result,ensure_ascii=False,indent=2)+"\n"
    if args.output:
        args.output.parent.mkdir(parents=True,exist_ok=True)
        args.output.write_text(rendered,encoding="utf-8")
    print(rendered)

if __name__=="__main__":
    main()
