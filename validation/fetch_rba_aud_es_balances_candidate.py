#!/usr/bin/env python3
import io
import json
import urllib.request
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

URL = 'https://www.rba.gov.au/statistics/tables/xls/a03.xlsx'
SERIES_ID = 'AESAT'
RUNTIME = Path('live_data/sections/NATIVE_LIQ_DATA.json')
OUT = Path('validation/liquidity-aud-es-balances-candidate.json')


def norm_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, str):
        s = v.strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d %b %Y', '%d %B %Y'):
            try:
                return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
            except ValueError:
                pass
    return None


def fetch_workbook():
    req = urllib.request.Request(URL, headers={'User-Agent': 'GMFQ-validation/1.0'})
    with urllib.request.urlopen(req, timeout=60) as r:
        body = r.read()
    if len(body) < 10000:
        raise SystemExit(f'RBA A3 workbook unexpectedly small: {len(body)} bytes')
    return load_workbook(io.BytesIO(body), data_only=True, read_only=False)


def runtime_series():
    rt = json.loads(RUNTIME.read_text(encoding='utf-8'))
    row = rt['AUD']['liquidity']
    out = {}
    for d, v in zip(row['dates'], row['values']):
        nd = norm_date(d)
        if not nd:
            raise SystemExit(f'unparseable runtime date {d!r}')
        out[nd] = float(v)
    return out


def official_series(wb):
    for ws in wb.worksheets:
        code_cell = None
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
            for cell in row:
                if str(cell.value or '').strip() == SERIES_ID:
                    code_cell = cell
                    break
            if code_cell:
                break
        if not code_cell:
            continue
        col = code_cell.column
        values = {}
        for r in range(code_cell.row + 1, ws.max_row + 1):
            val = ws.cell(r, col).value
            if not isinstance(val, (int, float)):
                continue
            nd = None
            for dc in range(1, min(col, 5) + 1):
                nd = norm_date(ws.cell(r, dc).value)
                if nd:
                    break
            if nd:
                values[nd] = float(val)
        if values:
            return ws.title, values
    raise SystemExit(f'RBA A3 series {SERIES_ID} not found')


def main():
    runtime = runtime_series()
    sheet, official = official_series(fetch_workbook())
    common = sorted(set(runtime) & set(official))
    exact_dates = [d for d in common if abs(runtime[d] - official[d]) <= 1e-9]
    mismatches = [
        {'date': d, 'runtime': runtime[d], 'official': official[d], 'diff': runtime[d] - official[d]}
        for d in common if abs(runtime[d] - official[d]) > 1e-9
    ]
    if len(common) < 250:
        raise SystemExit(f'insufficient RBA/runtime overlap: {len(common)}')
    ratio = len(exact_dates) / len(common)
    if ratio < 0.98:
        raise SystemExit(f'RBA AESAT historical exact-match ratio below contract: {ratio:.4f}')
    runtime_last = max(runtime)
    if runtime_last not in official or abs(runtime[runtime_last] - official[runtime_last]) > 1e-9:
        raise SystemExit('latest runtime AUD ES balance does not match RBA AESAT exactly')
    recent = common[-20:]
    if any(abs(runtime[d] - official[d]) > 1e-9 for d in recent):
        raise SystemExit('recent 20-date AUD ES balance window is not exact')
    latest_official = max(official)
    receipt = {
        'schema': 'GMFQ_LIQUIDITY_AUD_ES_BALANCES_CANDIDATE_V1',
        'status': 'GREEN_SOURCE_CANDIDATE_ONLY',
        'source': {
            'publisher': 'Reserve Bank of Australia',
            'table': 'A3 — ES Balances and Repo Agreements',
            'series_id': SERIES_ID,
            'series_name': 'Total Exchange Settlement Balances',
            'frequency': 'DAILY',
            'unit': 'AUD million',
            'url': URL,
            'sheet': sheet,
        },
        'runtime': {
            'count': len(runtime),
            'last_date': runtime_last,
            'last_value': runtime[runtime_last],
        },
        'proof': {
            'overlap_count': len(common),
            'exact_match_count': len(exact_dates),
            'exact_match_ratio': ratio,
            'recent_20_exact': True,
            'mismatch_count': len(mismatches),
            'mismatches': mismatches,
            'runtime_last_matches_official': True,
        },
        'latest_date': latest_official,
        'latest_value': official[latest_official],
        'candidate_pipeline_ready': True,
        'publication_performed': False,
        'runtime_mutated': False,
    }
    OUT.write_text(json.dumps(receipt, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps(receipt, indent=2, ensure_ascii=False))

if __name__ == '__main__':
    main()
