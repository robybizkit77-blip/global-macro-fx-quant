#!/usr/bin/env python3
import io
import json
import urllib.request
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

URL = 'https://www.rba.gov.au/statistics/tables/xls/a03.xlsx'
RUNTIME = Path('live_data/sections/NATIVE_LIQ_DATA.json')


def norm_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, str):
        s = v.strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d %b %Y', '%d %B %Y'):
            try:
                return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
            except ValueError:
                pass
    return None


def fetch():
    req = urllib.request.Request(URL, headers={'User-Agent': 'GMFQ-validation/1.0'})
    with urllib.request.urlopen(req, timeout=60) as r:
        body = r.read()
    if len(body) < 10000:
        raise SystemExit(f'RBA A3 workbook unexpectedly small: {len(body)} bytes')
    return body


def runtime_series():
    rt = json.loads(RUNTIME.read_text(encoding='utf-8'))
    row = rt['AUD']['liquidity']
    out = {}
    for d, v in zip(row['dates'], row['values']):
        nd = norm_date(d)
        if not nd:
            raise SystemExit(f'unparseable runtime date {d!r}')
        out[nd] = float(v)
    return out


def extract_candidate_columns(wb):
    candidates = []
    for ws in wb.worksheets:
        max_header_row = min(ws.max_row, 30)
        # Search each column's header cells for ES / Exchange Settlement balance language.
        for col in range(1, ws.max_column + 1):
            parts = [str(ws.cell(r, col).value or '').strip() for r in range(1, max_header_row + 1)]
            header = ' | '.join(x for x in parts if x)
            low = header.lower()
            if not (('exchange settlement' in low or 'es balance' in low or 'es balances' in low) and 'balance' in low):
                continue
            values = {}
            for r in range(1, ws.max_row + 1):
                val = ws.cell(r, col).value
                if not isinstance(val, (int, float)):
                    continue
                nd = None
                # dates are usually in one of the left-most columns
                for dc in range(1, min(col, 5) + 1):
                    nd = norm_date(ws.cell(r, dc).value)
                    if nd:
                        break
                if nd:
                    values[nd] = float(val)
            if values:
                candidates.append({'sheet': ws.title, 'column': col, 'header': header, 'values': values})
    return candidates


def compare(runtime, official):
    common = sorted(set(runtime) & set(official))
    exact = 0
    rounded0 = 0
    mismatches = []
    for d in common:
        rv, ov = runtime[d], official[d]
        if abs(rv - ov) <= 1e-9:
            exact += 1
        if abs(rv - round(ov)) <= 1e-9:
            rounded0 += 1
        if abs(rv - ov) > 1e-9:
            mismatches.append({'date': d, 'runtime': rv, 'official': ov, 'diff': rv-ov})
    return {
        'overlap': len(common),
        'exact': exact,
        'rounded0_exact': rounded0,
        'first_common': common[0] if common else None,
        'last_common': common[-1] if common else None,
        'mismatches_tail': mismatches[-12:],
    }


def main():
    runtime = runtime_series()
    wb = load_workbook(io.BytesIO(fetch()), data_only=True, read_only=False)
    rows = []
    for c in extract_candidate_columns(wb):
        cmp = compare(runtime, c['values'])
        rows.append({
            'sheet': c['sheet'],
            'column': c['column'],
            'header': c['header'],
            **cmp,
            'runtime_last_date': max(runtime),
            'runtime_last_value': runtime[max(runtime)],
            'official_on_runtime_last_date': c['values'].get(max(runtime)),
            'official_latest_date': max(c['values']),
            'official_latest_value': c['values'][max(c['values'])],
        })
    rows.sort(key=lambda x: (x['exact'], x['rounded0_exact'], x['overlap']), reverse=True)
    out = {
        'schema': 'GMFQ_RBA_AUD_ES_BALANCES_DIAGNOSTIC_V1',
        'status': 'DIAGNOSTIC_ONLY',
        'source_url': URL,
        'runtime_count': len(runtime),
        'candidates': rows[:20],
    }
    print(json.dumps(out, indent=2, ensure_ascii=False))
    if not rows:
        raise SystemExit('no A3 ES-balance candidate columns found')

if __name__ == '__main__':
    main()
