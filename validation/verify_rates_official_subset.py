#!/usr/bin/env python3
from __future__ import annotations
import csv, datetime as dt, html, io, json, pathlib, re, sys, urllib.error, urllib.request
from openpyxl import load_workbook

ROOT=pathlib.Path(__file__).resolve().parents[1]
RATES=json.loads((ROOT/'live_data'/'sections'/'NATIVE_RATES_DATA.json').read_text())
UA='Mozilla/5.0 GMFQ-rates-validator/1.0'

SOURCES={
 'USD':{
   'url':'https://home.treasury.gov/resource-center/data-chart-center/interest-rates/TextView?field_tdr_date_value=2026&type=daily_treasury_yield_curve',
   'kind':'treasury','required':True
 },
 'EUR':{
   'url':'https://data-api.ecb.europa.eu/service/data/YC/B.U2.EUR.4F.G_N_A.SV_C_YM.SR_{tenor}?startPeriod=2026-09-25&format=csvdata',
   'kind':'ecb_pair','required':True
 },
 'CAD':{
   'url':'https://www.bankofcanada.ca/rates/interest-rates/canadian-bonds/',
   'kind':'boc','required':True
 },
 'NZD':{
   'url':'https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/b/b2/hb2-daily-close.xlsx',
   'kind':'rbnz_xlsx','required':False
 }
}

def fetch_bytes(url:str)->bytes:
    req=urllib.request.Request(url,headers={'User-Agent':UA,'Accept':'text/csv,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,text/html,*/*'})
    with urllib.request.urlopen(req,timeout=45) as r:
        return r.read()

def textify(raw:str)->str:
    s=re.sub(r'<script[\s\S]*?</script>',' ',raw,flags=re.I)
    s=re.sub(r'<style[\s\S]*?</style>',' ',s,flags=re.I)
    s=re.sub(r'<[^>]+>',' ',s)
    return re.sub(r'\s+',' ',html.unescape(s)).strip()

def parse_treasury(raw:str):
    rows=re.findall(r'<tr[^>]*>([\s\S]*?)</tr>',raw,flags=re.I)
    found=[]
    for row in rows:
        cells=[textify(x) for x in re.findall(r'<t[dh][^>]*>([\s\S]*?)</t[dh]>',row,flags=re.I)]
        if not cells or not re.fullmatch(r'\d{2}/\d{2}/2026',cells[0]): continue
        nums=[]
        for c in cells[1:]:
            try: nums.append(float(c))
            except: nums.append(None)
        tail=nums[-8:]
        if len(tail)==8 and all(v is not None for v in tail):
            mm,dd,yyyy=cells[0].split('/')
            found.append((f'{yyyy}-{mm}-{dd}',tail[1],tail[5]))
    if not found: raise ValueError('Treasury: no parsable daily rows')
    return max(found,key=lambda x:x[0])

def parse_ecb_csv(raw:bytes):
    txt=raw.decode('utf-8-sig','replace')
    reader=csv.DictReader(io.StringIO(txt))
    found=[]
    for row in reader:
        d=row.get('TIME_PERIOD') or row.get('TIME PERIOD') or row.get('time_period')
        v=row.get('OBS_VALUE') or row.get('OBS VALUE') or row.get('obs_value')
        if not d or v in (None,''): continue
        try: found.append((str(d)[:10],float(v)))
        except: continue
    if not found: raise ValueError('ECB: no parsable TIME_PERIOD/OBS_VALUE rows')
    return max(found,key=lambda x:x[0])

def parse_ecb_pair(url_template:str):
    d2,y2=parse_ecb_csv(fetch_bytes(url_template.format(tenor='2Y')))
    d10,y10=parse_ecb_csv(fetch_bytes(url_template.format(tenor='10Y')))
    if d2!=d10: raise ValueError(f'ECB tenor dates differ: 2Y={d2},10Y={d10}')
    return d2,y2,y10

def parse_boc(raw:str):
    t=textify(raw)
    dates=sorted(set(x.replace('‑','-') for x in re.findall(r'2026[-‑]\d{2}[-‑]\d{2}',t)))
    if not dates: raise ValueError('BoC: no dates found')
    latest=dates[-1]
    exp2=float(RATES['CAD']['2Y']); exp10=float(RATES['CAD']['10Y'])
    if not re.search(r'\b'+re.escape(f'{exp2:.2f}')+r'\b',t): raise ValueError('BoC expected 2Y not present')
    if not re.search(r'\b'+re.escape(f'{exp10:.2f}')+r'\b',t): raise ValueError('BoC expected 10Y not present')
    return latest,exp2,exp10

def norm_header(v):
    return re.sub(r'\s+',' ',str(v or '').strip().lower())

def parse_rbnz_xlsx(raw:bytes):
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    candidates=[]
    for ws in wb.worksheets:
        rows=list(ws.iter_rows(values_only=True))
        header_idx=None; date_col=None; y2_col=None; y10_col=None
        for i,row in enumerate(rows[:40]):
            hs=[norm_header(v) for v in row]
            for j,h in enumerate(hs):
                if h=='date': date_col=j
                if h in {'2 year','2-year','2 year bond','2 year government bond'}: y2_col=j
                if h in {'10 year','10-year','10 year bond','10 year government bond'}: y10_col=j
            if date_col is not None and y2_col is not None and y10_col is not None:
                header_idx=i; break
        if header_idx is None: continue
        for row in rows[header_idx+1:]:
            if max(date_col,y2_col,y10_col)>=len(row): continue
            d=row[date_col]; y2=row[y2_col]; y10=row[y10_col]
            if isinstance(d,dt.datetime): d=d.date()
            if isinstance(d,dt.date) and isinstance(y2,(int,float)) and isinstance(y10,(int,float)):
                candidates.append((d.isoformat(),float(y2),float(y10)))
    if not candidates: raise ValueError('RBNZ workbook: no parsable Date/2 year/10 year rows')
    return max(candidates,key=lambda x:x[0])

def main()->int:
    results={}; failures=[]; blocked=[]; verified=[]
    for c,s in SOURCES.items():
        try:
            if s['kind']=='ecb_pair':
                date,y2,y10=parse_ecb_pair(s['url'])
            else:
                raw=fetch_bytes(s['url'])
                if s['kind']=='treasury': date,y2,y10=parse_treasury(raw.decode('utf-8','replace'))
                elif s['kind']=='boc': date,y2,y10=parse_boc(raw.decode('utf-8','replace'))
                elif s['kind']=='rbnz_xlsx': date,y2,y10=parse_rbnz_xlsx(raw)
                else: raise ValueError('unknown source kind')
            cur=RATES[c]
            row={
              'verification_status':'VERIFIED',
              'official_latest':date,'official_2Y':y2,'official_10Y':y10,
              'runtime_date':cur.get('date'),'runtime_2Y':cur.get('2Y'),'runtime_10Y':cur.get('10Y'),
              'date_relation':'SAME' if date==cur.get('date') else ('NEWER' if date>str(cur.get('date')) else 'OLDER'),
              'values_match': abs(float(y2)-float(cur.get('2Y')))<1e-9 and abs(float(y10)-float(cur.get('10Y')))<1e-9,
              'source':s['url']
            }
            results[c]=row; verified.append(c)
            if row['date_relation']=='OLDER': failures.append(f'{c}: official source parsed older than runtime')
            if row['date_relation']=='SAME' and not row['values_match']: failures.append(f'{c}: same-date official values mismatch runtime')
        except urllib.error.HTTPError as e:
            if not s.get('required') and e.code in (401,403):
                results[c]={'verification_status':'SOURCE_ACCESS_BLOCKED','http_status':e.code,'source':s['url']}
                blocked.append(c)
            else:
                results[c]={'verification_status':'ERROR','error':str(e),'source':s['url']}
                failures.append(f'{c}: {e}')
        except Exception as e:
            results[c]={'verification_status':'ERROR','error':str(e),'source':s['url']}
            failures.append(f'{c}: {e}')
    status='PASS_WITH_BLOCKED_SOURCES' if not failures and blocked else ('PASS' if not failures else 'FAIL')
    print(json.dumps({'status':status,'failures':failures,'verified':verified,'blocked':blocked,'results':results},indent=2))
    return 0 if not failures else 2

if __name__=='__main__': sys.exit(main())
