#!/usr/bin/env python3
"""Generate a read-only, fail-closed official CHF/NZD Rates staging candidate.

Staging NEVER modifies NATIVE_RATES_DATA, payload, manifest, scoring, or signals.
CHF SNB: source fetched from the pinned SNB URL for live evidence.
NZD RBNZ: official B2 XLSX can be reviewed offline with exact SHA and explicit
government-bond tenor mapping. Its output ALWAYS remains REVIEW_REQUIRED.
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import hashlib
import io
import json
import math
import re
import urllib.parse
from pathlib import Path

from build_rates_candidate_from_audit import append_hist, classify, reference_point
from probe_chf_nzd_rates_sources_v1 import (
    SNB_CSV, RBNZ_B2, fetch, parse_snb_csv
)

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "live_data/sections/NATIVE_RATES_DATA.json"
SCOPE = {"CHF", "NZD"}
BASIS = {
    "CHF": "Swiss National Bank · CHF government zero-coupon spot rates 2J/10J",
    "NZD": "Reserve Bank of New Zealand · B2 government benchmark close 2 year/10 year",
}
DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def checked_date(value):
    if not isinstance(value, str) or not DATE.fullmatch(value):
        raise ValueError("invalid ISO observation date")
    return dt.date.fromisoformat(value)


def numeric(v):
    if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v):
        raise ValueError("missing/non-finite numeric official yield")
    return float(v)


def parse_rbnz_b2_xlsx(raw: bytes, mapping: dict) -> dict:
    """Strict explicitly reviewed column map; never guess bond columns from bank bills/swaps."""
    if not raw.startswith(b"PK\x03\x04"):
        raise ValueError("RBNZ source is not an XLSX archive")
    if mapping.get("source_url") != RBNZ_B2:
        raise ValueError("official RBNZ B2 source identity mismatch")
    if mapping.get("source_sha256") != hashlib.sha256(raw).hexdigest():
        raise ValueError("official XLSX SHA256 differs from reviewed evidence")
    fields = ("sheet", "header_row", "group_row", "date_col", "y2_col", "y10_col")
    if not all(k in mapping for k in fields):
        raise ValueError("incomplete reviewed B2 worksheet/column map")
    hr, gr = mapping["header_row"], mapping["group_row"]
    dc, c2, c10 = (mapping[k] for k in ("date_col", "y2_col", "y10_col"))
    if any(type(n) is not int or n < 1 for n in (hr, gr, dc, c2, c10)):
        raise ValueError("worksheet coordinates must be positive integers")
    if len({dc, c2, c10}) != 3 or not (gr < hr and hr - gr <= 8):
        raise ValueError("government close tenor/header mapping ambiguous")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        if mapping["sheet"] not in wb.sheetnames:
            raise ValueError("reviewed worksheet not present")
        ws = wb[mapping["sheet"]]
        if ws.max_row < hr + 1:
            raise ValueError("no dated observations after reviewed B2 headers")
        group_values = [
            str(ws.cell(row=gr, column=c).value or "").lower()
            for c in range(1, max(c2, c10) + 1)
        ]
        if not any("government bond" in x and ("closing" in x or "close" in x)
                   for x in group_values):
            raise ValueError("government bond CLOSE section label not verified")
        def norm(x):
            return re.sub(r"\s+", " ", str(x or "").strip().lower())
        if norm(ws.cell(row=hr, column=dc).value) != "date":
            raise ValueError("Date header mismatch")
        if norm(ws.cell(row=hr, column=c2).value) not in {"2 year", "2 years"}:
            raise ValueError("2-year government bond header mismatch")
        if norm(ws.cell(row=hr, column=c10).value) not in {"10 year", "10 years"}:
            raise ValueError("10-year government bond header mismatch")
        pairs = {}
        for row in ws.iter_rows(min_row=hr + 1, max_col=max(dc,c2,c10), values_only=True):
            observed = row[dc - 1]
            if isinstance(observed, dt.datetime):
                day = observed.date()
            elif isinstance(observed, dt.date):
                day = observed
            elif isinstance(observed, str) and DATE.fullmatch(observed):
                day = dt.date.fromisoformat(observed)
            else:
                continue
            values = row[c2-1], row[c10-1]
            if any(v is None for v in values):
                continue
            y2, y10 = map(numeric, values)
            if not (0 <= y2 <= 18 and 0 <= y10 <= 18):
                raise ValueError("RBNZ government yield outside bounds")
            iso = day.isoformat()
            if iso in pairs and pairs[iso] != (y2, y10):
                raise ValueError("conflicting same-date RBNZ government yields")
            pairs[iso] = (y2, y10)
        if not pairs:
            raise ValueError("no same-day 2Y/10Y RBNZ government closing yields")
        iso = max(pairs)
        return {"date": iso, "2Y": pairs[iso][0], "10Y": pairs[iso][1],
                "observations":len(pairs), "basis": BASIS["NZD"]}
    finally:
        wb.close()


def validate_existing(cur: dict, currency: str):
    if currency not in SCOPE or currency not in cur:
        raise ValueError("unexpected currency")
    old = cur[currency]
    day = checked_date(old["date"])
    for tenor, hist in (("2Y","history2"),("10Y","history10")):
        h = old.get(hist, {})
        ds, vs = h.get("dates"), h.get("values")
        if not isinstance(ds,list) or not isinstance(vs,list) or not ds or len(ds)!=len(vs):
            raise ValueError("corrupt existing official Rates history")
        if ds!=sorted(set(ds)) or ds[-1]!=old["date"]:
            raise ValueError("last official history timestamp does not match snapshot")
        if abs(numeric(vs[-1])-numeric(old[tenor]))>1e-8:
            raise ValueError("official history tail/value mismatch")
        if not any(t.get("tenor")==tenor and t.get("date")==old["date"] and
                   abs(numeric(t.get("value"))-numeric(old[tenor]))<1e-8
                   for t in old.get("tenors",[])):
            raise ValueError("official current tenor snapshot mismatch")
    return day


def make_candidate(cur: dict, currency: str, observation: dict, *,
                   source_sha256: str, mode: str,
                   today: dt.date | None = None) -> tuple[dict | None,dict]:
    now = today or dt.datetime.now(dt.timezone.utc).date()
    previous = validate_existing(cur, currency)
    day = checked_date(observation["date"])
    y2, y10 = numeric(observation["2Y"]), numeric(observation["10Y"])
    lower, upper = (-3, 18) if currency=="CHF" else (0, 18)
    if not (lower <= y2 <= upper and lower <= y10 <= upper):
        raise ValueError("official yield outside conservative plausibility range")
    if day > now or (now-day).days > 5:
        raise ValueError("official observation is future/stale; no promotion")
    if day < previous:
        raise ValueError("new official observation predates canonical snapshot")
    old=cur[currency]
    receipt = {
        "schema":"GMFQ_CHF_NZD_OFFICIAL_STAGING_V1",
        "currency":currency,
        "source": SNB_CSV if currency=="CHF" else RBNZ_B2,
        "source_sha256":source_sha256,
        "method":mode,
        "official_basis":BASIS[currency],
        "as_of":day.isoformat(),
        "old_as_of":previous.isoformat(),
        "writes_live_data":False,
        "eligible_for_automatic_promotion":False,
        "review":"REQUIRED",
    }
    if day == previous:
        if abs(old["2Y"]-y2)>1e-8 or abs(old["10Y"]-y10)>1e-8:
            raise ValueError("official same-date observation conflicts with current canonical Rates")
        receipt["status"]="NO_NEW_OFFICIAL_PAIR"
        return None,receipt
    gap=(day-previous).days
    if gap>35:
        raise ValueError("official snapshot gap >35 days; must repair history first")
    max_jump_bp=min(250, 70 + gap * 12)
    for tenor,val in (("2Y",y2),("10Y",y10)):
        if abs(val-numeric(old[tenor]))*100>max_jump_bp:
            raise ValueError("unverified exceptional official yield jump; manual review mandatory")
    result=copy.deepcopy(cur)
    x=result[currency]
    label=BASIS[currency]
    (rd2,v2),cutoff=reference_point(x,"history2",day.isoformat())
    (rd10,v10),_=reference_point(x,"history10",day.isoformat())
    ch2=round((y2-v2)*100,1)
    ch10=round((y10-v10)*100,1)
    x.update(date=day.isoformat(),source=f"{label} · {day.isoformat()}",
             quality="FULL_CURRENT_OFFICIAL_STAGING_REVIEW_REQUIRED",
             **{"2Y":y2,"10Y":y10},
             curve_bp=round((y10-y2)*100,1),
             chg2_bp=ch2,chg10_bp=ch10,curve_state=classify(ch2,ch10),
             freshness_status="OFFICIAL_STAGING_REVIEW_REQUIRED",
             freshness_note="Not operational until independent source and runtime review.")
    for t in x.get("tenors",[]):
        if t["tenor"] in ("2Y","10Y"):
            t["value"]=y2 if t["tenor"]=="2Y" else y10
            t["date"]=day.isoformat()
    append_hist(x,"history2",day.isoformat(),y2)
    append_hist(x,"history10",day.isoformat(),y10)
    for suffix,refdate,refval,change in (("2",rd2,v2,ch2),("10",rd10,v10,ch10)):
        x["weekly"+suffix+"_validation"]={
            "start_date":refdate,"start_value":refval,
            "end_date":day.isoformat(),
            "end_value":y2 if suffix=="2" else y10,
            "change_bp":change,"source":x["source"],
            "note":f"Same-basis reference observed on/before {cutoff}; staging only."}
    changed=[k for k in result if result[k]!=cur[k]]
    if changed != [currency]:
        raise AssertionError(f"one-currency staging scope violated: {changed}")
    for tenor,h in (("2Y","history2"),("10Y","history10")):
        if (result[currency][h]["dates"][-1]!=day.isoformat() or
            abs(result[currency][h]["values"][-1]-result[currency][tenor])>1e-8):
            raise AssertionError("candidate history tail mismatch")
    receipt.update(status="GATED_CANDIDATE_REVIEW_REQUIRED",
                   y2=y2,y10=y10,change_from_snapshot_bp={
                       "2Y":round((y2-old["2Y"])*100,2),
                       "10Y":round((y10-old["10Y"])*100,2)},
                   changed_currencies=changed,
                   retained_other_currencies=True)
    return result, receipt


def collect(args) -> tuple[dict,bytes,str]:
    if args.currency=="CHF":
        if args.rbnz_mapping:
            raise ValueError("NZD mapping may not be used with CHF")
        if args.live:
            params=urllib.parse.urlencode({
                "dimSel":"D0(CHF),D1(2J,10J)","fromDate":"2026-09-01"})
            raw,_=fetch(SNB_CSV+"?"+params,accept="text/csv,text/plain")
            mode="PINNED_SNB_LIVE_FETCH"
        else:
            if not args.source_file:
                raise ValueError("SNB needs --live or an explicitly provided source file")
            raw=Path(args.source_file).read_bytes()
            mode="REVIEWED_LOCAL_SNB_SOURCE_UNAUTHENTICATED"
        p=parse_snb_csv(raw)
        obs={"date":p["as_of"],"2Y":p["2Y"],"10Y":p["10Y"]}
    else:
        if args.live or not args.source_file or not args.rbnz_mapping:
            raise ValueError("RBNZ B2 requires a reviewed official XLSX and exact SHA worksheet mapping")
        raw=Path(args.source_file).read_bytes()
        mapping=json.loads(Path(args.rbnz_mapping).read_text(encoding="utf8"))
        obs=parse_rbnz_b2_xlsx(raw,mapping)
        mode="REVIEWED_LOCAL_RBNZ_WORKBOOK_MAPPING"
    return obs,raw,mode


def main():
    p=argparse.ArgumentParser()
    p.add_argument("--currency",choices=sorted(SCOPE),required=True)
    p.add_argument("--live",action="store_true")
    p.add_argument("--source-file")
    p.add_argument("--rbnz-mapping")
    p.add_argument("--rates-file",default=str(BASE))
    p.add_argument("--output",required=True,help="Write ONLY to isolated staging destination")
    p.add_argument("--receipt",required=True)
    args=p.parse_args()
    if args.live and args.source_file:
        p.error("--live and --source-file are mutually exclusive")
    for rawpath in (args.output,args.receipt):
        target=Path(rawpath).resolve()
        if ROOT in target.parents or target==ROOT:
            raise SystemExit("stage artifacts cannot be written anywhere inside the tracked repo")
    cur=json.loads(Path(args.rates_file).read_text(encoding="utf8"))
    obs,raw,mode=collect(args)
    candidate,receipt=make_candidate(cur,args.currency,obs,
        source_sha256=hashlib.sha256(raw).hexdigest(),mode=mode)
    Path(args.receipt).parent.mkdir(parents=True,exist_ok=True)
    Path(args.receipt).write_text(json.dumps(receipt,indent=2,ensure_ascii=False)+"\n",encoding="utf8")
    if candidate is not None:
        Path(args.output).parent.mkdir(parents=True,exist_ok=True)
        Path(args.output).write_text(json.dumps(candidate,separators=(",",":"),ensure_ascii=False)+"\n",encoding="utf8")
    print(json.dumps(receipt,ensure_ascii=False))
    print("CHF_NZD_STAGING_ONLY_NO_LIVE_WRITE_PASS")


if __name__=="__main__":
    main()
