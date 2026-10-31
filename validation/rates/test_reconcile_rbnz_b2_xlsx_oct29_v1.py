#!/usr/bin/env python3
"""RBNZ XLSX publication comparison, synthetic workbook only; no data writes."""
from __future__ import annotations
import copy
import io
import json
import unittest
from pathlib import Path
from openpyxl import Workbook
from reconcile_rbnz_b2_xlsx_oct29_v1 import (
    ROOT,HTML,RATES,reconcile,locate_government_pair,
)
SOURCE=json.loads(HTML.read_text())
RATES_OBJ=json.loads(RATES.read_text())

def workbook(variant=None):
    book=Workbook()
    sheet=book.active;sheet.title="B2 Daily close"
    sheet.cell(1,1,"RBNZ B2 official research fixture (NOT REAL XLSX)")
    sheet.cell(4,6,"Secondary market government bond closing yields (%pa)")
    sheet.cell(4,9,"Swap rate spread close (bps)")
    heads={1:"Date",2:"OCR",3:"90 days",4:"1 year",
           6:"1 year",7:"2 year",8:"10 year",9:"2-10s"}
    for col,h in heads.items():sheet.cell(5,col,h)
    for i,row in enumerate(SOURCE["website_rows"],start=6):
        import datetime
        sheet.cell(i,1,datetime.date.fromisoformat(row["date"]))
        sheet.cell(i,6,4.0)
        sheet.cell(i,7,row["2Y"])
        sheet.cell(i,8,row["10Y"])
        sheet.cell(i,9,62.)
    if variant:variant(sheet)
    bio=io.BytesIO();book.save(bio);book.close()
    return bio.getvalue()

class IndependentWorkbookTests(unittest.TestCase):
    def test_exact_html_xlsx_ten_rows_match_no_auto_promotion(self):
        result=reconcile(workbook(),SOURCE,RATES_OBJ)
        self.assertEqual(result["status"],"OFFICIAL_HTML_XLSX_AGREE_REVIEWABLE_NOT_PUBLISHED")
        self.assertEqual(result["matched_count"],10)
        self.assertEqual(result["latest_pair_date"],"2026-10-29")
        self.assertEqual(result["latest_2Y"],3.77)
        self.assertEqual(result["latest_10Y"],4.87)
        self.assertEqual(result["columns_1_based"],{
            "date":1,"2Y_government_close":7,"10Y_government_close":8})
        self.assertFalse(result["eligible_for_automatic_promotion"])
        self.assertFalse(result["writes_operational_rates"])
        self.assertTrue(result["same_source_two_formats_not_independent_institution"])
    def test_untrusted_workbook_rejected(self):
        with self.assertRaises(ValueError):
            reconcile(b'not XLSX',SOURCE,RATES_OBJ)
    def test_swap_2_10s_never_government_10y(self):
        # No government group, only swap headings => fail closed.
        with self.assertRaises(ValueError):
            reconcile(workbook(lambda s:s.cell(4,6,"Swap rate spread close (bps)")),
                      SOURCE,RATES_OBJ)
    def test_false_government_10y_rejected(self):
        with self.assertRaisesRegex(ValueError,"mismatches"):
            reconcile(workbook(lambda s:s.cell(15,8,4.91)),SOURCE,RATES_OBJ)
    def test_false_government_2y_rejected(self):
        with self.assertRaisesRegex(ValueError,"mismatches"):
            reconcile(workbook(lambda s:s.cell(15,7,3.81)),SOURCE,RATES_OBJ)
    def test_source_release_governance_tampering_rejected(self):
        changed=copy.deepcopy(SOURCE)
        changed["governance"]["eligible_for_official_canonical_promotion"]=True
        with self.assertRaises(ValueError):
            reconcile(workbook(),changed,RATES_OBJ)
    def test_forged_canonical_pit_anchor_rejected(self):
        changed=copy.deepcopy(RATES_OBJ);changed["NZD"]["2Y"]+=0.01
        with self.assertRaises(ValueError):
            reconcile(workbook(),SOURCE,changed)
    def test_duplicate_government_column_pair_rejected(self):
        def mutate(s):
            s.cell(5,10,"2 year")
            s.cell(5,11,"10 year")
            # Group label right side also qualifies: two independent claimants.
            s.cell(4,10,"Secondary market government bond closing yields (%pa)")
        with self.assertRaises(ValueError):
            reconcile(workbook(mutate),SOURCE,RATES_OBJ)
    def test_newer_missing_date_not_filled_from_earlier(self):
        with self.assertRaisesRegex(ValueError,"missing HTML-publication date"):
            reconcile(workbook(lambda s:s.cell(15,1,None)),SOURCE,RATES_OBJ)

if __name__=="__main__":
    unittest.main(verbosity=2)
